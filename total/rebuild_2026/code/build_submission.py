"""Assemble the journal submission package in ZEB1/submission.

Inputs: polished manuscript Markdown, Supplementary Information Markdown,
finalized figures (finalize_journal_figures.R) and supplementary tables.
Outputs: manuscript DOCX (legends only and with embedded figures),
Supplementary Information DOCX, Supplementary Tables XLSX, cover letter DOCX.
"""
from pathlib import Path
import csv
import re
import shutil
import subprocess

import openpyxl
from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from openpyxl.styles import Alignment, Font

CODE = Path(__file__).resolve().parent
REBUILD = CODE.parent
ROOT = REBUILD.parents[1]
MS_DIR = REBUILD / "manuscript"
TABLES = REBUILD / "data/source_data_rebuilt/supplementary_tables"
CLIN_SENS = REBUILD / "data/source_data_rebuilt/revision2_statistics/clinical_age_sex_wbc_sensitivity.tsv"
SUB = ROOT / "submission"
FIG = SUB / "figures"
OUT_MS = SUB / "manuscript"
OUT_SUPP = SUB / "supplementary"
OUT_CL = SUB / "cover_letter"
BUILD = SUB / "_build"
for d in (OUT_MS, OUT_SUPP, OUT_CL, BUILD):
    d.mkdir(parents=True, exist_ok=True)

PANDOC = shutil.which("pandoc")
if not PANDOC:
    raise SystemExit("pandoc not found")


def superscripts(text: str) -> str:
    return re.sub(r"<sup>(.*?)</sup>", lambda m: "^" + m.group(1).replace(" ", "\\ ") + "^", text)


def reference_docx() -> Path:
    ref = BUILD / "reference.docx"
    subprocess.run([PANDOC, "-o", str(ref), "--print-default-data-file", "reference.docx"],
                   check=True)
    doc = Document(ref)
    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if name not in [s.name for s in doc.styles]:
            continue
        st = doc.styles[name]
        st.font.name = "Times New Roman"
        st.font.size = Pt(12)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        pf = st.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        pf.space_after = Pt(0)
        pf.space_before = Pt(0)
    for name, size in (("Title", 16), ("Heading 1", 14), ("Heading 2", 13), ("Heading 3", 12)):
        if name not in [s.name for s in doc.styles]:
            continue
        st = doc.styles[name]
        st.font.name = "Arial"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = None
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        st.paragraph_format.space_before = Pt(12)
        st.paragraph_format.space_after = Pt(6)
    if "Caption" in [s.name for s in doc.styles]:
        doc.styles["Caption"].font.italic = False
    doc.save(ref)
    return ref


def add_line_numbers(path: Path) -> None:
    doc = Document(path)
    for section in doc.sections:
        sect = section._sectPr
        ln = OxmlElement("w:lnNumType")
        ln.set(qn("w:countBy"), "1")
        ln.set(qn("w:restart"), "continuous")
        sect.append(ln)
        footer = section.footer.paragraphs[0]
        footer.alignment = 1
        run = footer.add_run()
        for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
            if tag:
                fld = OxmlElement("w:fldChar")
                fld.set(qn("w:fldCharType"), tag)
                run._r.append(fld)
            else:
                instr = OxmlElement("w:instrText")
                instr.set(qn("xml:space"), "preserve")
                instr.text = text
                run._r.append(instr)
    doc.save(path)


def to_docx(md_text: str, target: Path, ref: Path, line_numbers: bool) -> None:
    src = BUILD / (target.stem + ".md")
    src.write_text(md_text, encoding="utf-8")
    subprocess.run([PANDOC, str(src), "-f", "markdown+pipe_tables+superscript-raw_html",
                    "-t", "docx", "--reference-doc", str(ref),
                    "--resource-path", str(SUB), "-o", str(target)], check=True)
    if line_numbers:
        add_line_numbers(target)
    doc = Document(target)
    for i, para in enumerate(doc.paragraphs):
        if para._p.xpath(".//*[local-name()='oMathPara']"):
            para.paragraph_format.keep_together = True
            if i:
                doc.paragraphs[i - 1].paragraph_format.keep_with_next = True
    doc.save(target)
    print(f"  {target.relative_to(ROOT)}  ({target.stat().st_size // 1024} KB)")


def word_counts(md: str) -> tuple[int, int]:
    abstract = md.split("## Abstract", 1)[1].split("\n## ", 1)[0]
    body = md.split("## Introduction", 1)[1].split("## Data and code availability", 1)[0]
    clean = lambda s: len(re.sub(r"[#*\[\]()]", " ", s).split())
    return clean(abstract), clean(body)


# ── Main manuscript ─────────────────────────────────────────────────────────
main_md = (MS_DIR / "ZEB1_ZEB2_TALL_manuscript_polished.md").read_text(encoding="utf-8")
main_md = superscripts(main_md)
n_abs, n_body = word_counts(main_md)
n_refs = len(re.findall(r"^\d+\. ", main_md.split("## References", 1)[1], flags=re.M))
counts = (f"**Word count:** abstract {n_abs}; main text (Introduction to Discussion) {n_body}.  \n"
          f"**Figures:** 7. **References:** {n_refs}. **Supplementary material:** "
          f"9 figures and 8 tables; 1 main table.\n")
main_md = main_md.replace("\n---\n\n## Abstract", "\n" + counts + "\n---\n\n## Abstract", 1)

ref = reference_docx()
print("[manuscript]")
to_docx(main_md, OUT_MS / "ZEB1_ZEB2_TALL_manuscript.docx", ref, line_numbers=True)

with_figs = main_md
for i in range(1, 8):
    pat = re.compile(rf"(\*\*Figure {i}\..*?)(\n\n)", re.S)
    with_figs = pat.sub(lambda m, i=i: m.group(1) + f"\n\n![](figures/main/Figure{i}.png){{width=6.5in}}\n\n",
                        with_figs, count=1)
to_docx(with_figs, OUT_MS / "ZEB1_ZEB2_TALL_manuscript_with_figures.docx", ref, line_numbers=False)
(OUT_MS / "ZEB1_ZEB2_TALL_manuscript.md").write_text(main_md, encoding="utf-8")

# ── Supplementary Information ───────────────────────────────────────────────
si = (MS_DIR / "Supplementary_Information.md").read_text(encoding="utf-8")
for i in range(1, 10):
    pat = re.compile(rf"(## Supplementary Figure S{i}\..*?\n\n.*?)(\n\n|\Z)", re.S)
    images = (
        "\n\n".join(
            f"![](figures/supplementary/FigureS{i}_page{page}.png){{width=6.5in}}"
            for page in (1, 2)
        )
        if i <= 3 else f"![](figures/supplementary/FigureS{i}.png){{width=6.5in}}"
    )
    si = pat.sub(lambda m, i=i: m.group(1) + f"\n\n{images}\n\n",
                 si, count=1)
si = superscripts(si)
print("[supplementary]")
to_docx(si, OUT_SUPP / "ZEB1_ZEB2_TALL_Supplementary_Information_editorial_rebuild.docx", ref, line_numbers=False)

# ── Supplementary tables workbook ───────────────────────────────────────────
wb = openpyxl.Workbook()
wb.remove(wb.active)
sheets = [
    ("Table S1", "Clinical models and subtype attenuation", TABLES / "Table_S1_clinical_models.tsv"),
    ("Table S2", "Datasets, biological units and evidentiary roles", TABLES / "Table_S2_dataset_units.tsv"),
    ("Table S3", "Full 17-subtype developmental residual summary", TABLES / "Table_S4_full_subtype_residuals.tsv"),
    ("Table S4", "Clinical sensitivity adjusted for age, sex and white-cell count", CLIN_SENS),
]
drop_cols = {"source", "source_path", "model_id", "source_file"}
for name, title, path in sheets:
    with path.open(encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.reader(fh, delimiter="\t"))
    keep = [i for i, h in enumerate(rows[0]) if h.strip().lower() not in drop_cols]
    ws = wb.create_sheet(name)
    ws.append([f"Supplementary {name}. {title}"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([])
    for r, row in enumerate(rows):
        vals = []
        for i in keep:
            v = row[i] if i < len(row) else ""
            try:
                v = float(v) if re.fullmatch(r"-?\d+(\.\d+)?([eE][+-]?\d+)?", v) else v
            except ValueError:
                pass
            vals.append(v)
        ws.append(vals)
        if r == 0:
            for c in ws[ws.max_row]:
                c.font = Font(bold=True)
                c.alignment = Alignment(wrap_text=True, vertical="top")
    for col in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col[2:])
        ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 48)
    ws.freeze_panes = "A4"
xlsx = BUILD / "Supplementary_Source_Tables_S1-S4.xlsx"
wb.save(xlsx)
print(f"  {xlsx.relative_to(ROOT)}")

# ── Cover letter ────────────────────────────────────────────────────────────
cover = """To the Editor-in-Chief, *Haematologica*

Dear Editor,

We are pleased to submit our manuscript entitled **"Subtype-associated ZEB1–ZEB2 expression patterns in T-cell acute lymphoblastic leukemia after developmental-expression adjustment"** for consideration as an Original Article in *Haematologica*.

Contemporary genomic studies now resolve T-cell acute lymphoblastic leukemia (T-ALL) into at least 17 driver-defined molecular subtypes. We asked whether their relative ZEB1–ZEB2 expression differs after conditioning on a developmental-expression proxy. We combined four normal human thymus references with a 1,309-patient molecularly annotated diagnostic cohort and examined narrower contrasts in independent bulk RNA, patient-matched single-cell, lesion-defined and chromatin-contact data.

We believe three findings will interest your readership:

1. A ZEB-independent developmental-expression proxy accounts for 26% of ZEB1–ZEB2 balance variation in this cohort; molecular subtype adds 14 percentage points to the combined model (P = 7.8 × 10⁻⁴⁸). ETP-like leukemia, despite its low raw balance, lies closer to its cohort-fitted expectation at the same proxy score, whereas BCL11B, SPI1 and LMO2 γδ-like subtypes have more negative residuals.
2. All 12 BCL11B-rearranged leukemias in a lesion-defined series converge on ZEB2-dominant expression, including seven cases in which ZEB2 is not the rearrangement partner.
3. Unadjusted clinical associations of ZEB balance with induction failure and minimal residual disease largely attenuate after molecular-subtype adjustment. Fully adjusted Cox outputs are retained for audit only because of sparse-subtype convergence warnings; we do not claim independent clinical utility.

Within-cohort adjustment provides a common descriptive reference; a separate joint model tests conditional subtype association, and held-out subtype fits assess the reference curve's sensitivity to each target subtype. Alternative scores support the central BCL11B–ETP-like contrast within the diagnostic cohort.

The study used publicly available, de-identified data. The final analysis and figure code, the submitted supplementary table workbook, and aggregate robustness source summaries are public at https://github.com/JDjust/ZEB1-TALL-analysis. AI assistance with text, analysis-code preparation and data visualization is disclosed on the title page.

Thank you for considering our work.

Sincerely,

Chao Wu, on behalf of all authors
Department of Pulmonary Oncology, Tianjin Medical University Cancer Institute & Hospital, Tianjin, China
chaowutjmuch@163.com

Limei Li
Department of Hematology, The Second Affiliated Hospital of Hainan Medical University, Haikou, China
lilimei116@126.com
"""
print("[cover letter]")
to_docx(cover, OUT_CL / "Cover_letter_Haematologica.docx", ref, line_numbers=False)
print("Word counts: abstract", n_abs, "main text", n_body, "references", n_refs)
