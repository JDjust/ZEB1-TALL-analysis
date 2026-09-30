"""Audit the final submission package against its saved source records."""
from pathlib import Path
import json
import re
import unicodedata
from difflib import SequenceMatcher

import fitz
from docx import Document
import openpyxl

ROOT=Path(__file__).resolve().parents[3]
SUB=ROOT/"submission"
MS=ROOT/"total/rebuild_2026/manuscript/ZEB1_ZEB2_TALL_manuscript_polished.md"
MAN=ROOT/"total/rebuild_2026/data/manifests/references_verified.json"
OUT=SUB/"notes/EDITORIAL_REBUILD_INTEGRITY_AUDIT.md"

def norm(s):
    s=unicodedata.normalize("NFKD",s).lower()
    return re.sub(r"[^a-z0-9]+","",s)

s=MS.read_text(encoding="utf-8")
pre,refblock=s.split("## References\n\n",1)
raw=re.split(r"\n(?=\d+\.\s)",refblock.strip())
refs=[]
for x in raw:
    m=re.match(r"^(\d+)\.\s+(.+)",x,re.S)
    if m: refs.append((int(m.group(1)),m.group(2).strip()))
verified=json.loads(MAN.read_text(encoding="utf-8"))
rows=[]
for num,body in refs:
    direct=[v for v in verified if norm(v.get("title","")) in norm(body)]
    doi=re.search(r"doi:\s*([^\s]+)",body,re.I)
    if doi:
        direct=[v for v in verified if v.get("doi","").lower()==doi.group(1).rstrip(".").lower()] or direct
    if direct:
        v=direct[0]
        status="Crossref title/DOI match in saved manifest"
        score=1.0
    else:
        ranked=sorted(verified,key=lambda v:SequenceMatcher(None,norm(v.get("title","")),norm(body)).ratio(),reverse=True)
        v=ranked[0]
        status="REVIEW: no exact title/DOI match"
        score=SequenceMatcher(None,norm(v.get("title","")),norm(body)).ratio()
    rows.append((num,status,v.get("title",""),v.get("doi",""),score))

used=set()
for m in re.finditer(r"\[([0-9,\s–-]+)\]",pre):
    for item in m.group(1).split(","):
        item=item.strip()
        if not item: continue
        if "–" in item or "-" in item:
            nums=re.split(r"[–-]",item)
            if len(nums)==2 and all(x.strip().isdigit() for x in nums):
                used.update(range(int(nums[0]),int(nums[1])+1))
        elif item.isdigit(): used.add(int(item))

abs_text=s.split("## Abstract",1)[1].split("## Introduction",1)[0]
methods=s.split("## Methods",1)[1].split("## Data and code availability",1)[0]
body=s.split("## Introduction",1)[1].split("## Data and code availability",1)[0]
count=lambda t:len(re.sub(r"[#*\[\]()]"," ",t).split())
main_doc=Document(SUB/"manuscript/ZEB1_ZEB2_TALL_manuscript_with_figures.docx")
si_doc=Document(SUB/"supplementary/ZEB1_ZEB2_TALL_Supplementary_Information_editorial_rebuild.docx")
zh_doc=Document(SUB/"中文审阅版/正文_中文.docx")
zh_si_doc=Document(SUB/"中文审阅版/补充材料_中文.docx")
wb=openpyxl.load_workbook(SUB/"supplementary/Supplementary_Tables_S1-S8.xlsx",read_only=True)
main_pdf=fitz.open(SUB/"first_submission/Main_Article_with_Figures.pdf")
si_pdf=fitz.open(SUB/"first_submission/Supplementary_Information_complete.pdf")
checks={
    "abstract words":count(abs_text),
    "main text words":count(body),
    "Methods words":count(methods),
    "references":len(refs),
    "matched reference identities":sum("Crossref" in x[1] for x in rows),
    "main figures PDF":len(list((SUB/"figures/main").glob("Figure?.pdf"))),
    "supplementary figure numbers":len([i for i in range(1,10)
                                         if (SUB/f"figures/supplementary/FigureS{i}.pdf").exists()]),
    "main Word embedded figures":len(main_doc.inline_shapes),
    "supplement Word embedded pages":len(si_doc.inline_shapes),
    "Chinese main Word embedded figures":len(zh_doc.inline_shapes),
    "Chinese supplement Word embedded pages":len(zh_si_doc.inline_shapes),
    "main Word tables":len(main_doc.tables),
    "reader workbook sheets":len(wb.sheetnames),
    "main PDF pages":len(main_pdf),
    "supplement PDF pages":len(si_pdf),
    "main PDF contains Table 1":any("Table 1. Data resources" in p.get_text() for p in main_pdf),
    "supplement PDF contains S9":any("Supplementary Figure S9" in p.get_text() for p in si_pdf),
}
expected=list(range(1,len(refs)+1))
unused=sorted(set(expected)-used)
invalid=sorted(used-set(expected))
lines=["# Editorial rebuild integrity audit","",
       "Generated from the final manuscript, Crossref identity manifest, DOCX, XLSX and PDF files.","",
       "## Package checks","",
       "| Check | Value |","|---|---:|"]
lines += [f"| {k} | {v} |" for k,v in checks.items()]
lines += ["","## Reference use","",
          f"Final numbering consecutive: {([n for n,_ in refs]==expected)}.",
          f"Unused by numbered in-text citations: {unused or 'none'}.",
          f"Out-of-range citations: {invalid or 'none'}.",
          "Saved manifest records are source identity evidence; they do not establish that each citation supports the sentence where it appears.","",
          "| No. | Identity status | Title from saved Crossref record | DOI |",
          "|---:|---|---|---|"]
for n,status,title,doi,_ in rows:
    lines.append(f"| {n} | {status} | {title.replace('|','/')} | {doi} |")
lines += ["","## Required author checks","",
          "- Confirm author order, affiliations, contributions, funding, ethics and competing-interest wording.",
          "- Complete AI tool versions, access dates and human-verification details on the title page.",
          "- Review each reference in context; the title/DOI manifest checks identity only.",
          "- The public GitHub repository was synchronized with this rebuild on 2026-09-26; confirm the public files remain accessible before upload.",
          "- Independently review the target journal's live author guidelines before upload."]
OUT.parent.mkdir(parents=True,exist_ok=True)
OUT.write_text("\n".join(lines)+"\n",encoding="utf-8")
print(OUT)
print(checks)
print("unmatched refs",[(n,title) for n,status,title,_,_ in rows if "REVIEW" in status])
print("unused citations",unused,"invalid",invalid)
