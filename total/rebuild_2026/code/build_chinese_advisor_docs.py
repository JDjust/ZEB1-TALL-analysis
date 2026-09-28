"""Assemble Chinese reading copies with embedded, final-size journal plates."""
from pathlib import Path
import importlib.util
import subprocess
import tempfile

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm


SCRIPT = Path(__file__).resolve().with_name('build_editorial_journal_packages.py')
spec = importlib.util.spec_from_file_location('editorial_pkg', SCRIPT)
pkg = importlib.util.module_from_spec(spec)
spec.loader.exec_module(pkg)
OUT = pkg.SUB / 'article' / '00_Master'
OUT.mkdir(exist_ok=True)


def chinese_font(doc):
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3'):
        style = doc.styles[name]
        style.font.name = 'Microsoft YaHei'
        style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')


def convert(src, dest, images=()):
    subprocess.run([pkg.PANDOC, str(src), '-f', 'markdown+tex_math_dollars+superscript',
                    '-t', 'docx', '--reference-doc', str(pkg.REFERENCE_DOC),
                    '-o', str(dest)], check=True)
    doc = Document(dest)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.5)
    sec.top_margin = sec.bottom_margin = Cm(1.5)
    chinese_font(doc)
    for label, path in images:
        doc.add_page_break()
        p = doc.add_paragraph()
        p.add_run(label).bold = True
        img = doc.add_paragraph()
        img.alignment = 1
        img.add_run().add_picture(str(path), width=Cm(18))
    doc.save(dest)


SOURCE = pkg.BUILD
convert(SOURCE / '中文正文.md', OUT / '中文正文_含主图.docx',
        [(f'主图 Figure {i}', pkg.FIG / f'Figure{i}.png') for i in range(1, 8)])
convert(SOURCE / '中文补充图释.md', OUT / '中文补充图稿_含S1-S10.docx',
        [(f'补充图 Figure S{i}', pkg.SFIG / f'FigureS{i}.png') for i in range(1, 11)])
convert(SOURCE / '给老师讲的文章思路.md', OUT / '给老师讲的文章思路.docx')
print('Chinese advisor documents assembled:', [p.name for p in OUT.glob('*.docx')])
