"""Convert authoritative Python PDFs to native PowerPoint paths and text."""
from pathlib import Path
import argparse, json, hashlib, math, re
import fitz
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR
from pptx.shapes.autoshape import Shape
ROOT=Path(__file__).resolve().parents[2];OUT=ROOT/'03_publication_outputs/04_figures_optimized'
def sha(p):
 with Path(p).open('rb') as f:return hashlib.file_digest(f,'sha256').hexdigest()
def rgb(c):return RGBColor(*[int(round(max(0,min(1,v))*255)) for v in c[:3]])
def alpha(color,opacity):
 if opacity<1:
  node=color._color._xClr;el=OxmlElement('a:alpha');el.set('val',str(int(opacity*100000)));node.append(el)
def clip_line(p0,p1,r):
 # Liang-Barsky intersects original segments; independent coordinate clamping
 # would change oblique segments and Kaplan-Meier steps at the axes boundary.
 dx,dy=p1.x-p0.x,p1.y-p0.y;low,high=0.,1.
 for p,q in [(-dx,p0.x-r.x0),(dx,r.x1-p0.x),(-dy,p0.y-r.y0),(dy,r.y1-p0.y)]:
  if abs(p)<1e-12:
   if q<0:return None
  elif p<0:low=max(low,q/p)
  else:high=min(high,q/p)
  if low>high:return None
 return fitz.Point(p0.x+low*dx,p0.y+low*dy),fitz.Point(p0.x+high*dx,p0.y+high*dy)
def clip_polygon(points,r):
 # Only partially clipped filled marker curves require flattening. At 32
 # subdivisions per original cubic the deviation is below 0.004 PDF points.
 for axis,bound,keep in [(0,r.x0,1),(0,r.x1,-1),(1,r.y0,1),(1,r.y1,-1)]:
  if not points:break
  output=[];prior=points[-1];inside=(prior[axis]-bound)*keep>=-1e-8
  for current in points:
   now=(current[axis]-bound)*keep>=-1e-8
   if now!=inside:
    fraction=(bound-prior[axis])/(current[axis]-prior[axis])
    output.append(fitz.Point(prior.x+fraction*(current.x-prior.x),prior.y+fraction*(current.y-prior.y)))
   if now:output.append(current)
   prior=current;inside=now
  points=output
 return points
def source_clipped_drawings(page):
 result=[];stack=[];clips={'linear_paths':0,'filled_marker_paths':0,'removed_paths':0}
 for raw in page.get_drawings(extended=True):
  level=raw.get('level',0);stack=[v for v in stack if v[0]<level]
  if raw['type']=='clip':stack.append((level,raw['scissor']));continue
  if raw['type']=='group':continue
  d=dict(raw)
  if stack:
   r=fitz.Rect(stack[-1][1]);box=d['rect']
   outside=box.x0<r.x0-.01 or box.y0<r.y0-.01 or box.x1>r.x1+.01 or box.y1>r.y1+.01
   if outside:
    commands={item[0] for item in d['items']}
    if commands=={'l'} and d.get('fill') is None:
     items=[]
     for item in d['items']:
      segment=clip_line(item[1],item[2],r)
      if segment:items.append(('l',*segment))
     d['items']=items;clips['linear_paths']+=1
    elif commands=={'c'} and d.get('fill') is not None:
     points=[d['items'][0][1]]
     for item in d['items']:
      p0,p1,p2,p3=item[1:]
      for j in range(1,33):
       t=j/32.;u=1-t
       points.append(fitz.Point(u**3*p0.x+3*u*u*t*p1.x+3*u*t*t*p2.x+t**3*p3.x,u**3*p0.y+3*u*u*t*p1.y+3*u*t*t*p2.y+t**3*p3.y))
     points=clip_polygon(points,r)
     d['items']=[('l',points[j-1],point) for j,point in enumerate(points)] if points else []
     d['closePath']=True;clips['filled_marker_paths']+=1
    else:raise ValueError('Unsupported clipped native path '+str(commands))
    if not d['items']:clips['removed_paths']+=1;continue
    points=[p for item in d['items'] for p in item[1:] if isinstance(p,fitz.Point)]
    d['rect']=fitz.Rect(min(p.x for p in points),min(p.y for p in points),max(p.x for p in points),max(p.y for p in points))
  result.append(d)
 return result,clips
def native_path(slide,d):
 rect=d['rect'];x,y=rect.x0,rect.y0;w=max(rect.width,.01);h=max(rect.height,.01)
 # Direct native OOXML avoids python-pptx freeform's repeated whole-slide ID scan.
 sp=OxmlElement('p:sp');nv=OxmlElement('p:nvSpPr');cn=OxmlElement('p:cNvPr');cn.set('id',str(slide.shapes._next_shape_id));cn.set('name','Editable plot path');nv.append(cn);nv.append(OxmlElement('p:cNvSpPr'));nv.append(OxmlElement('p:nvPr'));sp.append(nv)
 spr=OxmlElement('p:spPr');sp.append(spr);xf=OxmlElement('a:xfrm');off=OxmlElement('a:off');off.set('x',str(int(Pt(x))));off.set('y',str(int(Pt(y))));ext=OxmlElement('a:ext');ext.set('cx',str(int(Pt(w))));ext.set('cy',str(int(Pt(h))));xf.append(off);xf.append(ext);spr.append(xf)
 geom=OxmlElement('a:custGeom');spr.append(geom)
 for tag in ['avLst','gdLst','ahLst','cxnLst']:geom.append(OxmlElement('a:'+tag))
 rec=OxmlElement('a:rect')
 for key,val in [('l','0'),('t','0'),('r','r'),('b','b')]:rec.set(key,val)
 geom.append(rec);pathlist=OxmlElement('a:pathLst');geom.append(pathlist)
 slide.shapes._spTree.insert_element_before(sp,'p:extLst');s=Shape(sp,slide.shapes)
 path=OxmlElement('a:path');path.set('w',str(int(w*12700)));path.set('h',str(int(h*12700)));pathlist.append(path)
 last=None
 def pt(parent,p):
  q=OxmlElement('a:pt');q.set('x',str(round((p.x-x)*12700)));q.set('y',str(round((p.y-y)*12700)));parent.append(q)
 def move(p):
  nonlocal last
  if last is None or abs(last.x-p.x)>1e-5 or abs(last.y-p.y)>1e-5:
   n=OxmlElement('a:moveTo');pt(n,p);path.append(n)
 def line(p):
  nonlocal last
  n=OxmlElement('a:lnTo');pt(n,p);path.append(n);last=p
 for item in d['items']:
  if item[0]=='l':move(item[1]);line(item[2])
  elif item[0]=='c':
   move(item[1]);n=OxmlElement('a:cubicBezTo')
   for p in item[2:]:pt(n,p)
   path.append(n);last=item[-1]
  elif item[0]=='re':
   r=item[1];ps=[r.tl,r.tr,r.br,r.bl];move(ps[0])
   for p in ps[1:]:line(p)
   path.append(OxmlElement('a:close'));last=None
  elif item[0]=='qu':
   q=item[1];ps=[q.ul,q.ur,q.lr,q.ll];move(ps[0])
   for p in ps[1:]:line(p)
   path.append(OxmlElement('a:close'));last=None
 if d.get('closePath'):path.append(OxmlElement('a:close'))
 if d.get('fill') is not None:
  s.fill.solid();s.fill.fore_color.rgb=rgb(d['fill']);alpha(s.fill.fore_color,d.get('fill_opacity',1))
 else:s.fill.background()
 if d.get('color') is not None and d.get('width',0)>0:
  s.line.color.rgb=rgb(d['color']);s.line.width=Pt(d['width']);alpha(s.line.color,d.get('stroke_opacity',1))
  pattern=re.match(r'\[([^\]]+)\]',d.get('dashes','') or '')
  if pattern:
   values=[float(value) for value in pattern.group(1).split()]
   if len(values)%2:values=values*2
   dash=OxmlElement('a:custDash')
   for j in range(0,len(values),2):
    part=OxmlElement('a:ds');part.set('d',str(max(1,round(values[j]/d['width']*100000))));part.set('sp',str(max(1,round(values[j+1]/d['width']*100000))));dash.append(part)
   spr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ln').append(dash)
 else:s.line.fill.background()
 s._element.spPr.append(OxmlElement('a:effectLst'))
 style=s._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style')
 if style is not None:s._element.remove(style)
 s.name='Editable plot path';return s
def native_text(slide,line,span):
 # Baseline and physical length retained; rotated axes remain native textboxes.
 x,y=span['origin'];fs=span['size'];dx,dy=line['dir'];angle=math.degrees(math.atan2(dy,dx))%360
 width=max(1,fitz.get_text_length(span['text'],fontname='hebo' if span['flags']&16 else 'helv',fontsize=fs)+fs*.4)
 height=fs*1.5
 if abs(angle-270)<1:
  left=x-height*.7;top=y-width;s=slide.shapes.add_textbox(Pt(left),Pt(top),Pt(height),Pt(width));s.rotation=270
  # Rotation around box centre: use an unrotated width x height box.
  s.left=Pt(x-width/2-height*.25);s.top=Pt(y-width/2-height*.25);s.width=Pt(width);s.height=Pt(height)
 else:
  cx=x+dx*width/2-dy*fs*.3;cy=y+dy*width/2+dx*(-fs*.3)
  s=slide.shapes.add_textbox(Pt(cx-width/2),Pt(cy-height/2),Pt(width),Pt(height));s.rotation=angle
 tf=s.text_frame;tf.clear();tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0;tf.word_wrap=False;tf.vertical_anchor=MSO_ANCHOR.TOP
 p=tf.paragraphs[0];p.space_before=p.space_after=Pt(0);p.line_spacing=1
 run=p.add_run();run.text=span['text'];run.font.name='Arial';run.font.size=Pt(fs);run.font.bold=bool(span['flags']&16);c=span['color'];run.font.color.rgb=RGBColor((c>>16)&255,(c>>8)&255,c&255)
 s.name='Editable text: '+span['text'][:40]
def image_with_original_alpha(block):
 # PDF RGB/indexed image bytes exclude their separate soft mask. Preserve that
 # mask when inserting a dense raster layer, so editable Office views retain
 # the authoritative PDF's point opacity and antialiased edges.
 if not block.get('mask'):return block['image']
 pix=fitz.Pixmap(block['image'])
 if pix.colorspace.n!=3:pix=fitz.Pixmap(fitz.csRGB,pix)
 return fitz.Pixmap(pix,fitz.Pixmap(block['mask'])).tobytes('png')
def source_paint_events(page,drawings,blocks):
 # Native label backgrounds can be painted after a raster atlas layer. Use
 # the PDF display-list indices, including each text baseline, in that case.
 log=page.get_bboxlog();image_log=[(i,fitz.Rect(item[1])) for i,item in enumerate(log) if item[0]=='fill-image']
 overlays=[d['seqno'] for d in drawings if d.get('fill')==(1.,1.,1.) and any(d['seqno']>seq and (d['rect'] & rect).get_area()>1 for seq,rect in image_log)]
 if not overlays:return None,[]
 traces=page.get_texttrace();events=[(d['seqno'],'draw',d) for d in drawings];prior=-1
 for block in blocks:
  if block['type']==1:
   rect=fitz.Rect(block['bbox']);matches=[seq for seq,box in image_log if max(abs(rect[j]-box[j]) for j in range(4))<.02]
   if len(matches)!=1:raise ValueError('Ambiguous PDF raster paint order')
   events.append((matches[0],'image',block))
  elif block['type']==0:
   for line in block['lines']:
    for span in line['spans']:
     matches=[t['seqno'] for t in traces if t['chars'] and abs(t['chars'][0][2][0]-span['origin'][0])<.02 and abs(t['chars'][0][2][1]-span['origin'][1])<.02 and ''.join(chr(char[0]) for char in t['chars'])==span['text']]
     if len(matches)!=1:
      if not span['text'].strip():seq=prior+.1
      else:raise ValueError('Ambiguous PDF text paint order: '+span['text'])
     else:seq=matches[0];prior=seq
     events.append((seq,'text',(line,span)))
 return sorted(events,key=lambda event:event[0]),overlays
def emit_picture(slide,block):
 from io import BytesIO
 r=block['bbox'];return slide.shapes.add_picture(BytesIO(image_with_original_alpha(block)),Pt(r[0]),Pt(r[1]),width=Pt(r[2]-r[0]),height=Pt(r[3]-r[1]))
def export(name):
 p=OUT/name/f'{name}.pdf';doc=fitz.open(p);page=doc[0];pr=Presentation();pr.slide_width=Pt(page.rect.width);pr.slide_height=Pt(page.rect.height);slide=pr.slides.add_slide(pr.slide_layouts[6]);slide.shapes.turbo_add_enabled=True
 drawings,clipping=source_clipped_drawings(page);texts=page.get_text('dict')['blocks']
 ordered,overlays=source_paint_events(page,drawings,texts)
 if ordered is not None:
  for _,kind,item in ordered:
   if kind=='draw':native_path(slide,item)
   elif kind=='text':native_text(slide,*item)
   else:emit_picture(slide,item)
 else:
  for i,d in enumerate(sorted(drawings,key=lambda drawing:drawing.get('seqno',0))):
   native_path(slide,d)
   if (i+1)%1000==0:print(name,f'{i+1}/{len(drawings)} paths converted',flush=True)
  for b in texts:
   if b['type']==0:
    for l in b['lines']:
     for s in l['spans']:native_text(slide,l,s)
   elif b['type']==1:emit_picture(slide,b)
 target=p.with_suffix('.pptx');pr.save(target)
 shapes=len(slide.shapes);text=sum(s.has_text_frame and bool(s.text) for s in slide.shapes)
 receipt={'figure':name,'pdf_sha256':sha(p),'pptx_sha256':sha(target),'shapes':shapes,'native_text_shapes':text,'native_drawing_shapes':len(drawings),'picture_shapes':sum(s.shape_type==13 for s in slide.shapes),'raster_soft_masks_preserved':True,'raster_soft_mask_layers':sum(b['type']==1 and bool(b.get('mask')) for b in texts),'source_rectangular_clipping':clipping,'native_dash_paths':sum(bool(re.match(r'\[([^\]]+)\]',d.get('dashes','') or '')) for d in drawings),'native_paths_and_text':True,'office_render_verified':False,'editing_verified':False,'code_sha256':sha(__file__)}
 receipt.update({'original_pdf_paint_order_for_raster_overlays':ordered is not None,'native_white_overlay_seqnos':overlays})
 (p.parent/'ppt_receipt.json').write_text(json.dumps(receipt,indent=2));print(name,shapes,'native objects',flush=True)
if __name__=='__main__':
 p=argparse.ArgumentParser();p.add_argument('figures',nargs='+');a=p.parse_args()
 for n in a.figures:export(n)
