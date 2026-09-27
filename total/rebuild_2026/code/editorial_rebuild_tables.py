"""Build one reader-facing S1–S10 workbook from frozen audit tables."""
from pathlib import Path
import math
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter
import editorial_rebuild_figures as m

SD=m.SD;D=m.D;read=m.read
OUT=m.OUT.parent/'Supplementary_Tables_S1-S10.xlsx'
wb=Workbook();wb.remove(wb.active)
navy='203342';blue='2B6687';white='FFFFFF';pale='F1F5F6';grey='697780'

def value(v):
    if pd.isna(v):return None
    if isinstance(v,(bool,)):return 'Yes' if v else 'No'
    if hasattr(v,'item'):
        try:v=v.item()
        except Exception:pass
    if isinstance(v,(int,float)) and not isinstance(v,bool):return v
    return str(v).replace('<U+03B3><U+03B4>','gd').replace('<U+03B1><U+03B2>','ab')

def sheet(name,title,note,df):
    ws=wb.create_sheet(name);ws.sheet_view.showGridLines=False
    df=df.copy().reset_index(drop=True)
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=max(2,len(df.columns)))
    ws.cell(1,1,title).font=Font(name='Arial',size=12,bold=True,color=white)
    ws.cell(1,1).fill=PatternFill('solid',fgColor=navy);ws.row_dimensions[1].height=24
    ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=max(2,len(df.columns)))
    ws.cell(2,1,note).font=Font(name='Arial',size=9,color=grey,italic=True)
    ws.cell(2,1).alignment=Alignment(wrap_text=True,vertical='center');ws.row_dimensions[2].height=34
    for j,col in enumerate(df.columns,1):
        c=ws.cell(4,j,str(col).replace('_',' '));c.fill=PatternFill('solid',fgColor='E5EBEE');c.font=Font(name='Arial',size=9,bold=True,color=navy)
        c.alignment=Alignment(wrap_text=True,vertical='center');c.border=Border(bottom=Side(style='thin',color='A3B0B7'))
    ws.row_dimensions[4].height=35
    for i,row in enumerate(df.itertuples(index=False,name=None),5):
        for j,x in enumerate(row,1):
            v=value(x);c=ws.cell(i,j,v);c.font=Font(name='Arial',size=9,color=navy)
            c.alignment=Alignment(vertical='top',wrap_text=True)
            if i%2==0:c.fill=PatternFill('solid',fgColor='F8FAFB')
            if isinstance(v,float):c.number_format='0.00E+00' if (v!=0 and abs(v)<.001) else '0.000' if abs(v)<1 else '0.00'
        ws.row_dimensions[i].height=16 if len(df.columns)<12 else 18
    for j,col in enumerate(df.columns,1):
        longest=max([len(str(col))]+[min(len(str(value(x) or '')),34) for x in df[col].head(300)])
        ws.column_dimensions[get_column_letter(j)].width=max(12,min(35,longest+3))
    ws.freeze_panes='A5';ws.auto_filter.ref=f'A4:{get_column_letter(len(df.columns))}{4+len(df)}'
    ws.print_title_rows='1:4';ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0;ws.sheet_properties.outlinePr.summaryBelow=True
    return ws

def main():
    rows=[
        ('GSE142522','Normal thymus','Sorted stage library','6 stage summaries','Normal stage shape','Independent normal resource'),
        ('GSE195812','Normal thymus','Pooled FACS stage library','8 stage summaries from six donors','Normal stage shape','Pooled units, not eight donors'),
        ('GSE206710','Normal thymus','Sample-stage aggregate','12 aggregates from three donors','Normal reference shape','Donor nesting retained'),
        ('Park/HTA','Normal thymus','Donor-stage aggregate','15 eligible paired donors','Normal stage shape','Independent reference'),
        ('Yayon pediatric CITE/HTSA','Normal thymus','Donor × author state','5 RNA donors; 33 eligible states','High-resolution state pattern','ETP observed in one RNA donor'),
        ('Yayon spatial CMA','Normal thymus','Spatial donor','6 Visium donors','Corticomedullary localization','Direct ZEB2 detection not evaluable'),
        ('Pölönen 2024','Pediatric/AYA T-ALL','Diagnostic patient','1,309; 17 subtypes','Discovery, residual, program freeze','Shared with ALLCatchR2 development'),
        ('AIEOP120','Pediatric T-ALL','Diagnostic patient','120; 16 official classes','External pediatric architecture','33 fusion-anchor exact matches'),
        ('GSE146901','ETP/non-ETP T-ALL','RNA library','8 ETP; 10 non-ETP; 4 peripheral T','Independent bulk program polarity','Peripheral healthy T is limited comparator'),
        ('GSE234608','ETP/MPAL and conventional T-ALL','RNA sample','11 ETP/MPAL; 17 conventional','Independent bulk program polarity','Assay-specific group definitions'),
        ('GSE243914','Blast-enriched T-ALL','RNA sample','6 ETP; 18 non-ETP','Blast-enriched program persistence','Not malignant-pure'),
        ('GSE248287','Single-cell T-ALL','Patient malignant pseudobulk','15 diagnosis patients','Author-malignant program preservation','Author label; no invented CNV/TCR'),
        ('GSE253355','Normal bone marrow','Donor × broad lineage','12 donors','Lineage-affinity audit','No cell-of-origin inference'),
        ('GSE162280','BCL11B-rearranged acute leukemia','Case RNA library','12 cases','Partner convergence','AML/MPAL/ETP mix; no ordinary T-ALL control'),
        ('GSE280250','Author adult T-ALL','Diagnostic RNA sample','79 cases; 72 aged ≥18','Adult architecture test','No listed overlap with ALLCatchR2 development cohorts'),
        ('ALLCatchR2','Cross-age subtype classifier','Official high-confidence call','79 GSE280250 classifications','Adult subtype context','Not an independent discovery test'),
        ('GSE165209 / GSE173432','CD34+ acute BCL11B OE','Paired perturbation unit','See S10 source rows','Reverse causal boundary','Acute OE does not induce frozen ZEB2-side'),
        ('HiChIP / pooled loops / scATAC','Chromatin','Library or peak set','3 ETP+4 non-ETP HiChIP; 7 scATAC sets','Exploratory context','Cross-assay polarity not reproduced'),
        ('Lim 2024','Single-cell T-ALL','Patient pseudobulk','54 Day0; 41 state pairs; 12 Day28 pairs','Within-patient state and null boundary','Distinct biological denominators'),
    ]
    sheet('S1','Table S1. Data resources and evidentiary roles','Counts are eligible biological units. Overlap language is limited to published cohort lists; patient identity was not independently verified.',pd.DataFrame(rows,columns=['Resource','Context','Independent unit','N','Evidentiary role','Boundary']))

    normal=read(D/'a4_yayon/a4a_stage_balance.tsv')[['stage','n_donors','n_cells','balance_median','dev_median','pole']]
    cma=read(D/'a4_yayon/a4ab_stage_balance_and_cma.tsv')[['stage','cma_weighted_median']]
    normal=normal.merge(cma,on='stage',how='left',validate='one_to_one')
    sheet('S2','Table S2. Normal-thymus author-state and CMA mapping','All 33 eligible RNA states are shown. The rows are descriptive and must not be treated as an ordered pseudotime path; locked early/DP/SP poles are defined separately.',normal)

    st=read(SD/'F3/F3D-F_subtype_statistics.tsv')
    st['Residual median (95% CI)']=st.apply(lambda r:f"{r.residual_within_median:.2f} ({r.residual_within_ci_low:.2f} to {r.residual_within_ci_high:.2f})",axis=1)
    st['Adjusted effect (95% CI)']=st.apply(lambda r:f"{r.adjusted_effect:.2f} ({r.adjusted_effect_ci_low:.2f} to {r.adjusted_effect_ci_high:.2f})",axis=1)
    sub=st[['subtype','n','dev_median','Residual median (95% CI)','Adjusted effect (95% CI)','adjusted_effect_hc3_p','adjusted_effect_bh_fdr','outside_normal']].rename(columns={'dev_median':'Median D','adjusted_effect_hc3_p':'HC3 P','adjusted_effect_bh_fdr':'BH FDR','outside_normal':'Outside measured normal D'})
    sheet('S3','Table S3. Seventeen-subtype developmental residuals','Patient median residual and joint-model subtype coefficient answer different questions; bootstrap CI is patient-stratified. Full-precision estimates remain in source TSV.',sub)

    s4=read(D/'aieop120_official/aieop120_subtype_residual_medians.tsv')
    sheet('S4','Table S4. AIEOP120 official classifier audit','Published argmax labels, no extra probability cutoff. The 33 fusion-anchor cases match exactly. External architecture support is moderate, not a 17-class replication.',s4)

    frozen=read(D/'gate2_program/frozen_ZEB_side50.tsv')
    a6=read(D/'a6_yayon/a6_2_gene_effects.tsv')[['symbol','delta_normal','class']]
    frozen=frozen.merge(a6,on='symbol',how='left',validate='one_to_one')
    frozen=frozen[['side','rank','symbol','gene_id','beta_R','partial_R2','adj.P.Val','delta_normal','class']]
    sheet('S5','Table S5. Frozen ZEB-side programs and normal-development classes','The 50 genes per side were frozen before downstream validation; 90/100 were evaluable in Yayon normal donor×stage data. Blank normal class means unrecovered.',frozen)

    pairs=read(SD/'revision3_robustness/bcl11b_etp_matched_pairs.tsv')
    sheet('S6a','Table S6a. Developmentally matched BCL11B–ETP-like pairs','18 one-to-one pairs without replacement; all values are primary-cohort units. This is a within-cohort contrast.',pairs)
    lesions=read(SD/'F4/F4C_GSE162280_cases.tsv')
    sheet('S6b','Table S6b. Complete BCL11B lesion-defined series','Twelve AML/MPAL/ETP cases; five ZEB2 partners and seven other partners. Ratios are descriptive; no ordinary T-ALL control is available.',lesions)

    extrows=[('GSE146901','bulk T-ALL',84,84,.869,.596),('GSE234608','bulk ETP/MPAL',90,90,.800,.640),('GSE243914','blast-enriched',98,98,.918,.770),('GSE248287','author malignant cells',94,94,.777,.620),('GSE280250','adult bulk',86,78,.987,.858)]
    ex=pd.DataFrame(extrows,columns=['Cohort','Assay context','Recovered frozen genes','Tested gene effects','Expected-direction fraction','Gene-effect Spearman'])
    a65=read(D/'a65_external/a65_summary.tsv');
    for cohort in ['GSE146901','GSE234608']:
        row=a65[(a65.cohort==cohort)&(a65['set']=='all_recovered')]
        if len(row):
            idx=ex.index[ex.Cohort==cohort][0];ex.loc[idx,'Recovered frozen genes']=int(row.iloc[0]['n']);ex.loc[idx,'Expected-direction fraction']=float(row.iloc[0]['expected_frac']);ex.loc[idx,'Gene-effect Spearman']=float(row.iloc[0]['rho_delta_vs_neg_beta'])
    sheet('S7','Table S7. External frozen-program preservation','Fractions apply to evaluable gene effects, with assay-specific recovered/tested denominators in archived TSVs. A directional structure is supported; each gene is not claimed to replicate independently.',ex)

    lineage=read(D/'a7_lineage/scores/a7a_unit_scores_min20.tsv')
    sheet('S8a','Table S8a. Normal marrow donor × lineage scores','Pseudobulk units with ≥20 cells; ≥10-cell sensitivity is archived. Side scores use recovered frozen genes only.',lineage)
    decomp=read(D/'a7c_decompose/a7c_zeb2_split.tsv')
    sheet('S8b','Table S8b. ZEB2-side marrow-affinity decomposition','45 myeloid-concordant genes and three non-myeloid genes; the three-gene slice is not a newly selected program.',decomp)

    malignant=read(D/'a7b_malignant/a7b_patient_scores.tsv')
    sheet('S9a','Table S9a. Author-malignant patient pseudobulks','Fifteen diagnostic patients; Celltypes_all==Malignant as supplied by authors. ZEB2-side score association is weak and its interval spans zero.',malignant)
    adult=read(D/'a8_adult/a8_patient_scores.tsv')
    sheet('S9b','Table S9b. Adult cohort patient-level balance, proxy and frozen scores','All 79 author Adult T-ALL diagnoses are primary. B, D and R are refitted within this adult cohort; seven patients are younger than 18.',adult)
    adultsub=read(D/'a8_adult/a8_2_summary.tsv')
    sheet('S9c','Table S9c. Official ALLCatchR2 subtype context','High-confidence calls are primary; candidate calls are sensitivity. Compound classes are not forced into a 17-class map.',adultsub)

    oe=read(D/'a3_gse165209/gse173432/a3_cameraPR_frozen_side50.tsv')
    sheet('S10a','Table S10a. Acute BCL11B-overexpression boundary','Observed CAMERA directions are opposite the prespecified ZEB-side state expectations; small P values do not support induction.',oe)
    chrom=read(SD/'S10/S10A-C_HiChIP_sample_contacts.tsv')
    sheet('S10b','Table S10b. HiChIP sample-level contacts','Three ETP and four non-ETP T-ALL libraries comprise the main exploratory comparison; exact two-sided P=.057 is the discrete minimum.',chrom)
    scatac=read(SD/'S10/S10D_scATAC_peak_sets.tsv')
    sheet('S10c','Table S10c. scATAC cross-assay boundary','Seven peak sets do not reproduce the HiChIP polarity; peak counts do not establish transcriptional regulation.',scatac)
    clin=read(SD/'S6/S6A_all_endpoint_models.tsv')
    sheet('S10d','Table S10d. Clinical attenuation and survival warnings','Clinical inference is exploratory. Sparse-subtype Cox convergence warnings prevent EFS/OS from supporting independent survival claims; FDR cannot repair unstable coefficients.',clin)
    ridge=read(SD/'revision4_targeted/ridge_cox_sensitivity.tsv')
    sheet('S10e','Table S10e. Penalized Cox sensitivity audit','Subtype dummy coefficients were penalized; balance, age, sex and WBC remained unpenalized. This does not establish an independent prognostic biomarker.',ridge)

    OUT.parent.mkdir(parents=True,exist_ok=True);wb.save(OUT)
    print(OUT, 'sheets=',len(wb.sheetnames),'numbered tables=',10)

if __name__=='__main__':main()
