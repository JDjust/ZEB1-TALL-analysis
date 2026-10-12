from pathlib import Path
import subprocess,json,re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from lxml import etree
from zipfile import ZipFile
B=Path(__file__).resolve().parent.parent;M=B/'manuscript';Q=B/'QA'
import shutil
exe=next((B/'reproduction_code/tools').glob('*/pandoc.exe'),None) or shutil.which('pandoc')
assert exe, 'Install Pandoc 3.12.1 or supply the recorded portable Windows executable'
style=Document();sec=style.sections[0]
sec.top_margin=sec.bottom_margin=Inches(.8);sec.left_margin=sec.right_margin=Inches(.85)
for name in ['Normal','Body Text','First Paragraph']:
    if name not in style.styles:style.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    st=style.styles[name];st.font.name='Cambria';st.font.size=Pt(11);st.font.color.rgb=RGBColor(0,0,0);st.paragraph_format.space_after=Pt(6);st.paragraph_format.line_spacing=2.0
for name,size in [('Title',20),('Heading 1',15),('Heading 2',13),('Heading 3',11)]:
    st=style.styles[name];st.font.name='Cambria';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
style.save(M/'DOCUMENT_REFERENCE_STYLE.docx')
reports=[]
for src,name in [('MANUSCRIPT_SOURCE.md','MANUSCRIPT'),('SUPPLEMENTARY_SOURCE.md','SUPPLEMENTARY_INFORMATION')]:
    common=[str(exe),str(M/src),'--standalone','--citeproc','--bibliography',str(B/'literature/REFERENCES.bib'),'-f','markdown+tex_math_dollars','--csl',str(B/'literature/vancouver.csl')]
    for ext,target,args in [('docx',name+'.docx',['--reference-doc',str(M/'DOCUMENT_REFERENCE_STYLE.docx')]),('html',name+'.html',['--math-method=mathml','--embed-resources']),('markdown',name+'.md',[])]:
        cp=subprocess.run(common+args+['-t',ext,'-o',str(M/target)],capture_output=True,text=True,encoding='utf-8');assert cp.returncode==0,(cp.stdout,cp.stderr)
        assert 'not found' not in cp.stderr.lower(),cp.stderr
        reports.append(dict(document=target,pandoc_returncode=cp.returncode,warnings=cp.stderr.strip()))
    p=M/(name+'.docx');doc=Document(p)
    footer=doc.sections[0].footer.paragraphs[0];footer.alignment=2
    run=footer.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');run._r.addnext(field)
    for section in doc.sections:
        lines=OxmlElement('w:lnNumType');lines.set(qn('w:countBy'),'1');lines.set(qn('w:start'),'1');lines.set(qn('w:restart'),'continuous');lines.set(qn('w:distance'),'240');section._sectPr.append(lines)
    for section in doc.sections:
        lines=OxmlElement('w:lnNumType');lines.set(qn('w:countBy'),'1');lines.set(qn('w:start'),'1');lines.set(qn('w:restart'),'continuous');lines.set(qn('w:distance'),'240');section._sectPr.append(lines)
    if name=='MANUSCRIPT':
        # Keep the full author/affiliation title page together before the abstract.
        for paragraph in doc.paragraphs:
            if paragraph.text.strip()=='Abstract':
                paragraph.paragraph_format.page_break_before=True
                break
            paragraph.paragraph_format.line_spacing=1.0
            paragraph.paragraph_format.space_after=Pt(5)
            if not paragraph.style.name.startswith('Heading'):
                for run in paragraph.runs:run.font.size=Pt(10.5)
    doc.save(p)
    with ZipFile(p) as z:
        x=etree.fromstring(z.read('word/document.xml'));ns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math','w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        eq=len(x.findall('.//m:oMath',ns));paras=[''.join(e.itertext()) for e in x.findall('.//w:t',ns)]
        assert not any('@Polonen' in v for v in paras)
        if name=='MANUSCRIPT':assert eq>=6,eq
        reports.append(dict(document=p.name,editable_equation_objects=eq,xml_readable=True,citations_resolved=True))
# Structured formula map for actual computational objects.
formulas=re.findall(r'\$\$\s*(.*?)\s*\$\$',(M/'04_methods.md').read_text(),re.S)
purposes=['Fixed normal-reference score','Separate Lim cycle/depth standardized rank-OLS estimand','Subtype log-expression residualization','AIIMS within-phenotype midrank-residual association','Landmark time-varying Cox ZEB1 block','Resampled apparent discrimination difference']
purposes=['Global average-tie rank standardization','Subtype-adjusted global-rank ZEB1 response model','Fixed normal-reference score','Fixed positive-member omission score','Full-library depth-corrected donor-state ZEB1','Lim cycle/depth-adjusted score response model','Subtype-residual expression for pathway ranking','AIIMS within-phenotype rank residual correlation','MRD landmark time-varying ZEB1 hazard block','Apparent nested Harrell discrimination increment']
purposes=['Global average-tie rank standardization','Subtype-adjusted global-rank ZEB1 response model','Fixed normal-reference score','Fixed positive-member omission score','Full-library depth-corrected donor-state ZEB1','Lim cycle/depth-adjusted score response model','Subtype-residual expression for pathway ranking','AIIMS within-phenotype rank residual correlation','MRD landmark time-varying ZEB1 hazard block','Apparent nested Harrell discrimination increment']
assert len(formulas)==len(purposes)==10==10
(M/'FORMULA_MAP.json').write_text(json.dumps([dict(equation_id='E'+str(i+1),scientific_purpose=p,latex=f,source='manuscript/04_methods.md',performed_implementation='METHOD_SOURCE_RECORDS.json; exact named model records in Methods',biological_unit='patient except normal member reference is donor',scope='Exact performed score/model; document export fits no model') for i,(p,f) in enumerate(zip(purposes,formulas))],indent=2),encoding='utf-8')
(Q/'DOCUMENT_ASSEMBLY_QA.json').write_text(json.dumps(reports,indent=2),encoding='utf-8')
print(json.dumps(reports,indent=2))
