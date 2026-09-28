"""Build four lean journal packages from the locked editorial manuscript and plates.

The source texts, plots and tables are already frozen. This script formats them,
checks journal word/display limits and assembles editable manuscript files.
"""
from pathlib import Path
import copy
import re
import shutil
import subprocess
import tempfile

import fitz
from docx import Document
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[3]
SUB = ROOT / 'submission'
BUILD = SUB / '_build' / 'editorial_rebuild'
PACK = SUB / 'article'
FIG = SUB / 'article' / '00_Figure_Master' / 'main'
SFIG = SUB / 'article' / '00_Figure_Master' / 'supplementary'
TABLES = BUILD / 'Supplementary_Tables_S1-S10.xlsx'
PANDOC = shutil.which('pandoc')
CHROME = Path(r'C:/Program Files/Google/Chrome/Application/chrome.exe')
REFERENCE_DOC = SUB / '_build' / 'reference.docx'

MASTER = (BUILD / 'Editorial_Master.md').read_text(encoding='utf-8')
LEGENDS = (BUILD / 'Editorial_Figure_Legends.md').read_text(encoding='utf-8')
SI_LEGENDS = (BUILD / 'Editorial_Supplement_Legends.md').read_text(encoding='utf-8')
BJH_BODY = (BUILD / 'BJH_Body.md').read_text(encoding='utf-8')

JOURNALS = {
    'npj Systems Biology and Applications': {
        'title': 'Developmental context separates subtype ZEB states in T cell acute lymphoblastic leukemia',
        'abstract': ('Normal thymopoiesis patterns ZEB1–ZEB2 expression, complicating its interpretation in T-cell acute lymphoblastic leukemia (T-ALL). '
                     'We integrated four normal-thymus references, a spatial atlas and 1,309 diagnostic patients across 17 molecular subtypes. '
                     'A ZEB-independent developmental-expression proxy explained 25.6% of balance variance; subtype added 13.9 percentage points. '
                     'BCL11B had the most negative within-cohort residual (median −1.76; 95% CI −2.47 to −0.98), while TLX3 was positive (+0.64). '
                     'Frozen residual-associated programs differed from normal early-to-double-positive changes: 13/90 genes were concordant, 32 discordant and 45 weak. '
                     'Eighteen developmentally matched BCL11B–ETP-like pairs differed in residual (mean −1.92; CI −2.98 to −0.67). '
                     'Directional program structure persisted across independent bulk, blast-enriched, author-defined malignant-cell and adult cohorts. '
                     'The results describe a subtype-associated, developmentally non-equivalent expression architecture without establishing cell of origin, mechanism or clinical utility.'),
        'abstract_limit': 150, 'body_limit': None,
    },
    'JCMM': {
        'title': 'Developmentally distinct ZEB programs across molecular subtypes of T-cell acute lymphoblastic leukemia',
        'abstract': ('The relative expression of ZEB1 and ZEB2 varies during human thymopoiesis, but its relationship to molecular T-cell acute lymphoblastic leukemia (T-ALL) subtypes is unclear. '
                     'We integrated normal thymus and spatial references with 1,309 diagnostic pediatric and young-adult patients across 17 subtypes. '
                     'After conditioning on a ZEB-independent developmental-expression proxy, BCL11B had the lowest median residual (−1.76; 95% CI −2.47 to −0.98) and TLX3 the highest (+0.64). '
                     'A frozen 100-gene residual-associated program differed from normal early-to-double-positive development: 13/90 evaluable genes were concordant, 32 discordant and 45 weak. '
                     'Eighteen BCL11B cases differed from developmental-score-matched ETP-like cases (mean paired residual difference −1.92; CI −2.98 to −0.67). '
                     'All 12 lesion-defined BCL11B cases were ZEB2-dominant, including seven without a ZEB2 partner. '
                     'Frozen program directions were retained across independent bulk, blast-enriched, malignant-cell and adult resources. '
                     'Normal myeloid affinity did not reduce the leukemia state to passive admixture alone. '
                     'These observational findings describe a developmentally non-equivalent molecular phenotype without proving direct regulation or clinical utility.'),
        'abstract_limit': 200, 'body_limit': 4500,
    },
    'British Journal of Haematology': {
        'title': 'ZEB expression states distinguish BCL11B from developmentally matched ETP-like T-ALL',
        'abstract': ('T-cell acute lymphoblastic leukaemia (T-ALL) subtypes may resemble different stages of thymopoiesis, complicating interpretation of ZEB1 and ZEB2 expression. '
                     'We analysed four normal thymus references, spatial context and 1,309 paediatric and young-adult diagnostic cases across 17 molecular subtypes. '
                     'A three-gene developmental-expression proxy explained 25.6% of ZEB balance variance; subtype added 13.9 percentage points. '
                     'BCL11B had the lowest within-cohort residual (median −1.76; 95% CI −2.47 to −0.98), whereas TLX3 was positive (+0.64). '
                     'The frozen residual-associated transcriptional programme differed from normal early-to-double-positive development. '
                     'Eighteen BCL11B cases had lower residual than matched ETP-like cases (mean paired difference −1.92; CI −2.98 to −0.67). '
                     'All 12 lesion-defined BCL11B cases were ZEB2-dominant, including seven without a ZEB2 partner. '
                     'Independent bulk, blast-enriched, malignant-cell and adult resources retained directional programme structure. '
                     'Normal myeloid affinity and specimen adjustment argued against passive admixture as its sole explanation. '
                     'The evidence supports a developmentally contextualised subtype phenotype, without establishing a direct mechanism or prognostic utility.'),
        'abstract_limit': 200, 'body_limit': 3000,
    },
    'Scientific Reports': {
        'title': 'Developmentally contextualized ZEB1–ZEB2 expression states in T-cell acute lymphoblastic leukemia',
        'abstract': ('ZEB1–ZEB2 balance changes during thymopoiesis, yet molecular T-cell acute lymphoblastic leukemia (T-ALL) subtypes may differ beyond that developmental pattern. '
                     'We integrated normal thymus and spatial resources with 1,309 pediatric and young-adult diagnostic cases in 17 subtypes. '
                     'A ZEB-independent developmental-expression proxy explained 25.6% of balance variance; adding subtype increased raw explained variance by 13.9 percentage points. '
                     'BCL11B showed the lowest median within-cohort residual (−1.76; 95% CI −2.47 to −0.98), while TLX3 was positive (+0.64). '
                     'Among 90 evaluable frozen residual-associated genes, 13 aligned with normal early-to-double-positive change, 32 opposed it and 45 changed weakly. '
                     'Eighteen BCL11B–ETP-like matched pairs differed in residual (mean −1.92; CI −2.98 to −0.67). '
                     'Independent bulk, blast-enriched, author-malignant and adult data retained directional program structure. '
                     'Normal myeloid affinity did not make passive admixture a sufficient explanation. '
                     'The observational data support a subtype-associated expression architecture; no full 17-subtype external replication, mechanism or clinical biomarker is claimed.'),
        'abstract_limit': 200, 'body_limit': 4500,
    },
}


def word_count(text):
    return len(re.findall(r"\b[\w]+(?:[–'-][\w]+)*\b", text, re.UNICODE))


def body_from_master():
    return MASTER.split('## Introduction\n', 1)[1].split('## Data and code availability\n', 1)[0]


def compact_body(body):
    # These are secondary contextual paragraphs; the seven result sections and
    # all primary models remain intact in the JCMM/SR article versions.
    drop_starts = [
        'Several types of evidence are needed to answer that question',
        'The spatial thymus data add anatomical context',
        'The adult analysis extends the directional architecture',
        'Perturbation, chromatin and outcomes place clear limits',
    ]
    blocks = body.split('\n\n')
    blocks = [x for x in blocks if not any(x.startswith(start) for start in drop_starts)]
    out = '\n\n'.join(blocks)
    out = out.replace('The analysis used the executable R/Python scripts, source-data TSVs, eligibility manifests and locked audit notes in the project archive. R 4.6.0 was used in the prior locked analysis; individual package versions are retained in the source session information and should be read from that file rather than inferred from this manuscript. New figure plates were rendered directly at final physical size with vector text. The exact figure-source mappings are recorded in `PANEL_SOURCE_MAP.md`. No new analysis family was introduced during the editorial rebuild.',
        'Executable R/Python scripts, source tables, eligibility manifests and session information are archived with the analysis. Figure plates were rendered at final physical size with vector text.')
    return out


def references():
    # Freeze the bibliography independently of archived prior manuscripts.
    frozen = (BUILD / 'Editorial_References.md').read_text(encoding='utf-8').strip()
    assert len(re.findall(r'^\d+\. ', frozen, flags=re.M)) == 52
    return frozen


def common_tail(journal):
    availability = ('Public datasets and their roles are listed by accession in Supplementary Table S1. '
        'Analysis scripts, frozen gene membership, the formatted supplementary workbook and non-identifying aggregate outputs are available at '
        'https://github.com/JDjust/ZEB1-TALL-analysis (tag: v2.0.4-print-qc). '
        'Original patient-level data retain their source access conditions.')
    contributor = MASTER.split('## Author contributions\n', 1)[1].split('## Figure legends\n', 1)[0]
    contributor = contributor.replace('## Artificial intelligence use\n\n', '## Artificial intelligence use\n\n')
    if journal in ('npj Systems Biology and Applications', 'Scientific Reports'):
        # Nature Portfolio requests Methods placement for LLM disclosure. The
        # truthful scope is retained; the responsible authors verify final text.
        contributor = contributor.split('## Artificial intelligence use\n',1)[0]
    return '## Data and code availability\n\n' + availability + '\n\n## Author contributions\n\n' + contributor


def make_markdown(journal):
    info = JOURNALS[journal]
    pre = MASTER.split('## Introduction\n', 1)[0]
    pre = re.sub(r'^# .*$', '# ' + info['title'], pre, count=1, flags=re.M)
    pre = re.sub(r'(?s)(## Abstract\n\n).*', lambda m:m.group(1)+info['abstract']+'\n\n', pre, count=1)
    if journal == 'British Journal of Haematology':
        body = BJH_BODY
    elif journal == 'npj Systems Biology and Applications':
        body = '## Introduction\n\n' + body_from_master()
    else:
        body = '## Introduction\n\n' + compact_body(body_from_master())
    if journal in ('npj Systems Biology and Applications', 'Scientific Reports'):
        ai = ('### Use of artificial intelligence tools\n\n'
              'OpenAI ChatGPT/Codex and Cursor assisted with language editing, code drafting/debugging, '
              'figure layout and document assembly. Authors executed the numerical analyses, inspected source tables '
              'and final figure plates, and retain responsibility for all results and wording.\n\n')
        body = body.replace('### Reproducibility and software\n\n', ai+'### Reproducibility and software\n\n')
        if journal == 'Scientific Reports' and ai not in body:
            body += '\n\n' + ai
    md = pre + body.strip() + '\n\n' + common_tail(journal).strip() + '\n\n## Figure legends\n\n' + LEGENDS.strip() + '\n\n## References\n\n' + references() + '\n'
    assert word_count(info['abstract']) <= info['abstract_limit'], (journal,word_count(info['abstract']))
    main = body.split('## Introduction\n',1)[1]
    assert info['body_limit'] is None or word_count(main) <= info['body_limit'], (journal,word_count(main))
    for number in range(1,8):
        assert f'Figure {number}.' in md
    return md,word_count(main),word_count(info['abstract'])


def make_master_markdown():
    """Assemble the full advisor-readable master with the current figures."""
    pre = MASTER.split('## Data and code availability\n', 1)[0].strip()
    return (pre + '\n\n' + common_tail('JCMM').strip() +
            '\n\n## Figure legends\n\n' + LEGENDS.strip() +
            '\n\n## References\n\n' + references() + '\n')


def pandoc_docx(markdown, dest, scratch):
    src = scratch / 'manuscript.md'
    src.write_text(markdown,encoding='utf-8')
    cmd = [PANDOC,str(src),'-f','markdown+tex_math_dollars+superscript','-t','docx',
           '--reference-doc',str(REFERENCE_DOC),'-o',str(dest)]
    subprocess.run(cmd,check=True)
    doc = Document(dest)
    sec = doc.sections[0]
    sec.page_width = Cm(21); sec.page_height = Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.5)
    sec.top_margin = sec.bottom_margin = Cm(1.5)
    for i in range(1,8):
        doc.add_page_break()
        p = doc.add_paragraph()
        p.add_run(f'Figure {i}').bold = True
        pic = doc.add_paragraph()
        pic.alignment = 1
        pic.add_run().add_picture(str(FIG / f'Figure{i}.png'),width=Cm(18))
    doc.save(dest)


SI_INTRO = '''# Supplementary Information

## Scope and structure

This Supplementary Information accompanies the seven main figures and the numbered workbook Tables S1–S10. Figures S1–S10 are each supplied as one physically sized, single-page plate in the second half of this PDF. The legends below are detailed audit descriptions; the full-precision values and denominator definitions remain in the workbook and source tables. The primary analysis, model equations and biological-unit definitions are in the main manuscript Methods. This document adds eligibility, cross-assay and negative-boundary details rather than a second discovery narrative.

## Supplementary methods and interpretation notes

Normal thymus harmonization retained each study's annotated stages and biological units. GSE195812's eight pooled libraries and GSE206710's 12 stage aggregates within three donors contribute 20 measured reference units, not 20 independent donors. Figure 1C displays all 33 eligible Yayon author states; its within-family plotting order is descriptive rather than inferred pseudotime. Spatial CMA localization uses author cell2location state assignments in six donors. Direct ZEB2 Visium spot detection was not considered evaluable.

The Pölönen 1,309-patient three-gene proxy and common residual are the primary descriptive reference. Leave-one-subtype-out fits change the training population for each class and therefore assess self-influence without replacing that common baseline. The leave-one-marker-out, expanded-score and spline degrees-of-freedom 2/3/4 sensitivities were fixed comparisons. Fourteen of 18 primary BCL11B cases exceed the measured normal proxy domain; normal projection remains a biological reference, not the residual estimand. The subtype median residual and joint-model coefficient have different estimands. Cross-validated model comparisons split at patient level.

The 100 frozen genes were selected only in the discovery cohort. External scores use measured genes and the locked membership, and Table S7 separates recovered from actually model-tested genes. The A6 normal-development comparison uses donor-by-stage pseudobulks, with 90 genes eligible. The matched BCL11B–ETP-like comparison remains observational and within cohort even when those classes are excluded from program derivation. The 12 lesion-defined cases are not an ordinary T-ALL control series.

For normal marrow, donor-by-lineage pseudobulks require at least 20 cells for the primary result and 10 for sensitivity. The specimen model compares 928 bone marrow and 381 blood patients; its coefficient concordance alone cannot establish cell-intrinsic signal. GSE243914 is blast-enriched, and GSE248287 uses source-author malignant labels only. Lim denominators are 54 Day-0 pseudobulks, 41 within-patient ZBTB16 state pairs and 12 Day-0/Day-28 pairs; these are not interchangeable. Adult GSE280250 recalculates all standardized quantities within its 79-case cohort and retains explicit age restrictions. The ALLCatchR2 calls are context on the same cohort, not a separate validation.

Acute BCL11B overexpression did not induce the expected frozen ZEB2-side direction. HiChIP had three ETP and four non-ETP usable libraries; the exact two-sided rank test gives P=0.057. Pooled loops and seven scATAC peak sets did not independently reproduce the polarity. Clinical binary results attenuate with subtype; sparse-subtype Cox warnings prevent inferential survival claims even though original output and ridge sensitivity are retained in Table S10. No cutoff, calibration or clinical utility was evaluated.

## Supplementary figure legends
'''


def make_supplement_md():
    return SI_INTRO + '\n' + SI_LEGENDS.strip() + '\n\n## Supplementary tables\n\n' + \
        'Tables S1–S10 are supplied in one formatted Excel workbook. Each numbered table may contain multiple tabs to preserve patient, gene and sensitivity denominators. Full precision remains in source TSVs.\n'


def chrome_pdf(markdown, dest, scratch):
    md = scratch / 'si_text.md'; html = scratch / 'si_text.html'
    md.write_text(markdown,encoding='utf-8')
    subprocess.run([PANDOC,str(md),'-f','markdown+tex_math_dollars','-t','html5','-s','--mathml','-o',str(html)],check=True)
    css = ("<style>@page{size:A4;margin:17mm 19mm 16mm 19mm}body{font:10.4pt/1.35 Arial;color:#172B39}"
           "h1,h2,h3{break-after:avoid;color:#172B39}h1{font-size:17pt}h2{font-size:12pt;margin-top:1.1em}"
           "p{orphans:3;widows:3}math[display='block']{display:block;text-align:center}</style>")
    html.write_text(html.read_text(encoding='utf-8').replace('</head>',css+'</head>',1),encoding='utf-8')
    with tempfile.TemporaryDirectory(prefix='zeb_si_',dir=scratch) as profile:
        subprocess.run([str(CHROME),'--headless=new','--disable-gpu','--no-first-run','--no-pdf-header-footer',
                        f'--user-data-dir={profile}',f'--print-to-pdf={dest}',html.as_uri()],
                       check=True,timeout=180,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)


def merge_supplement(text_pdf, dest):
    out=fitz.open()
    with fitz.open(text_pdf) as source: out.insert_pdf(source)
    for i in range(1,11):
        with fitz.open(SFIG/f'FigureS{i}.pdf') as plate:
            assert plate.page_count==1
            out.insert_pdf(plate)
    out.save(dest,garbage=4,deflate=True)
    out.close()


def cover_letter(journal,dest):
    info=JOURNALS[journal]
    doc=Document();sec=doc.sections[0]
    sec.left_margin=sec.right_margin=Cm(2.4)
    title=doc.add_paragraph('Cover letter')
    title.style='Title'
    doc.add_paragraph('Dear Editor,')
    doc.add_paragraph(f'Please consider our Original Article, “{info["title"]}”, for publication in {journal}.')
    if journal.startswith('npj'):
        pitch=('We use a developmental-context model and frozen transcriptional programs to separate '
               'normal thymic patterning from subtype-associated ZEB expression in T-ALL. The model is tested '
               'by held-out subtype fits, alternative expression proxies, external bulk and adult cohorts.')
    elif journal=='JCMM':
        pitch=('The biological focus is a BCL11B-associated ZEB2-skewed phenotype beyond matched ETP-like '
               'developmental-expression context. A frozen program and normal thymus comparison establish '
               'that the phenotype is not a simple continuation of early thymocyte expression.')
    elif journal=='British Journal of Haematology':
        pitch=('The hematologic focus is the distinction between an immature ETP-like transcriptional context '
               'and the additional ZEB2-side state in BCL11B-associated disease. The work integrates modern '
               'molecular subtypes with normal thymopoiesis and patient-level validation resources.')
    else:
        pitch=('The manuscript provides a technically auditable, public-data analysis of a bounded subtype-associated '
               'expression claim, with fixed program selection, patient-level units, sensitivity checks and explicit '
               'negative boundaries for causal, chromatin and clinical interpretations.')
    doc.add_paragraph(pitch)
    doc.add_paragraph('The primary cohort contains 1,309 diagnostic patients across 17 molecular subtypes. '
        'The revised submission includes seven main figures, ten single-page supplementary figures and one '
        'formatted workbook of Tables S1–S10. The main article presents observational associations only; '
        'acute perturbation, chromatin and clinical results are limited according to their evidentiary strength.')
    doc.add_paragraph('All authors have reviewed the manuscript. The work uses public de-identified data; '
        'competing-interest, funding, ethics, data and AI-use statements are included in the manuscript. '
        'The corresponding authors will complete any portal-specific declarations.')
    if journal=='JCMM':
        old=Document(PACK/'JCMM'/'Cover_Letter.docx')
        pubs=[p.text.strip() for p in old.paragraphs if 'doi:' in p.text.lower()]
        assert len(pubs)>=2
        doc.add_paragraph('Recent relevant publications by corresponding author Chao Wu (past 3–5 years):')
        for p in pubs[:4]:doc.add_paragraph(p,style='List Bullet')
    doc.add_paragraph('Sincerely,\nChao Wu and Limei Li\nCorresponding authors\nchaowutjmuch@163.com; lilimei116@126.com')
    doc.save(dest)


def build():
    assert PANDOC and CHROME.exists() and REFERENCE_DOC.exists() and TABLES.exists()
    with tempfile.TemporaryDirectory(prefix='zeb_editorial_',dir=BUILD) as tmp:
        scratch=Path(tmp)
        si_text=scratch/'SI_text.pdf'
        chrome_pdf(make_supplement_md(),si_text,scratch)
        for journal in JOURNALS:
            folder=PACK/journal;folder.mkdir(parents=True,exist_ok=True)
            md,body_n,abstract_n=make_markdown(journal)
            journal_scratch=scratch/re.sub(r'[^A-Za-z0-9]','_',journal)
            journal_scratch.mkdir()
            pandoc_docx(md,folder/'Manuscript.docx',journal_scratch)
            merge_supplement(si_text,folder/'Supplementary_Information.pdf')
            shutil.copy2(TABLES,folder/'Supplementary_Tables_S1-S10.xlsx')
            figures=folder/'Figures';figures.mkdir(exist_ok=True)
            for i in range(1,8):shutil.copy2(FIG/f'Figure{i}.pdf',figures/f'Figure{i}.pdf')
            cover_letter(journal,folder/'Cover_Letter.docx')
            print(journal,'abstract',abstract_n,'main',body_n,'SI pages',fitz.open(folder/'Supplementary_Information.pdf').page_count)


if __name__=='__main__':
    build()
