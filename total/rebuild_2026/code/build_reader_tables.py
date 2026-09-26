"""Build reader-facing tables from frozen, full-precision source TSV files."""
from __future__ import annotations

import csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[3]
DATA = ROOT / "total/rebuild_2026/data/source_data_rebuilt"
GATE = ROOT / "total/data/validation/gateA_lim2025"
OUT = ROOT / "submission/supplementary/Supplementary_Tables_S1-S8.xlsx"
MAIN = ROOT / "submission/tables"
MAIN.mkdir(parents=True, exist_ok=True)


def read(path):
    with path.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f, delimiter="\t"))


def fnum(x, digits=2):
    if x is None or x == "" or str(x).upper() == "NA":
        return "—"
    return f"{float(x):.{digits}f}"


def ci(r, key="estimate", lo="ci_low", hi="ci_high"):
    return f"{fnum(r[key])} ({fnum(r[lo])}–{fnum(r[hi])})"


def pnum(x):
    if x is None or x == "" or str(x).upper() == "NA":
        return "—"
    v = float(x)
    return f"{v:.2e}" if v < .001 else f"{v:.3g}"


book = Workbook()
book.remove(book.active)


def add_sheet(name, title, headers, rows, note="", audit_survival=False):
    ws = book.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws.append([title])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws.row_dimensions[1].height = 25
    ws["A1"].font = Font(name="Arial", size=12, bold=True, color="203744")
    ws["A1"].alignment = Alignment(vertical="center")
    ws.append([note])
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    ws["A2"].font = Font(name="Arial", size=9, italic=True, color="52646C")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 39 if len(note) > 200 else (30 if len(note) > 110 else 23)
    ws.append(headers)
    ws.row_dimensions[3].height = 34
    edge = Side(style="hair", color="CBD5D8")
    for c in ws[3]:
        c.fill = PatternFill("solid", fgColor="244455")
        c.font = Font(name="Arial", size=9, bold=True, color="FFFFFF")
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = Border(bottom=edge)
    for i, row in enumerate(rows, 4):
        ws.append(row)
        ws.row_dimensions[i].height = 42 if any(len(str(v)) > 95 for v in row) else (28 if any(len(str(v)) > 55 for v in row) else 22)
        for cell in ws[i]:
            cell.font = Font(name="Arial", size=9, color="203744")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.border = Border(bottom=edge)
            if i % 2:
                cell.fill = PatternFill("solid", fgColor="F4F7F8")
            if audit_survival and row[0] in ("Event-free survival", "Overall survival"):
                cell.font = Font(name="Arial", size=9, italic=True, color="6D777B")
                cell.fill = PatternFill("solid", fgColor="EDF0F1")
    for j in range(1, len(headers)+1):
        vals = [str(headers[j-1])] + [str(r[j-1]) for r in rows]
        width = min(44, max(12, max(min(len(v), 46) for v in vals) + 2))
        ws.column_dimensions[get_column_letter(j)].width = width
    ws.freeze_panes = "B4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(headers))}{ws.max_row}"
    ws.print_title_rows = "1:3"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.orientation = "landscape"
    return ws


raw_clin = read(DATA/"supplementary_tables/Table_S1_clinical_models.tsv")
full_clin = {r["endpoint"]:r for r in read(DATA/"revision2_statistics/clinical_age_sex_wbc_sensitivity.tsv")}
endpoints = list(dict.fromkeys(r["endpoint"] for r in raw_clin))
clin = {(r["endpoint"], r["adjustment"]):r for r in raw_clin}
s1 = []
s4 = []
for ep in endpoints:
    a = clin[(ep,"Unadjusted")]
    b = next(r for r in raw_clin if r["endpoint"]==ep and r["adjustment"]!="Unadjusted")
    c = full_clin[ep]
    n_ev = f'{a["n"]}/{a["events"]}'
    s1.append([ep, a["measure"], n_ev, ci(a), pnum(a["p"]), pnum(a["endpoint_bh_fdr"]),
               ci(b), pnum(b["p"]), pnum(b["endpoint_bh_fdr"]), ci(c), pnum(c["p"]), pnum(c["endpoint_bh_fdr"])])
    s4.append([ep, c["method"], f'{c["n"]}/{c["events"]}', ci(c),
               pnum(c["p"]), pnum(c["endpoint_bh_fdr"]), pnum(c["ph_balance_p"]),
               pnum(c["ph_global_p"]),
               ("AUDIT ONLY — " if c["method"] == "Cox PH" else "") + c["status"]])
add_sheet("Table S1", "Supplementary Table S1 | Clinical attenuation across three models",
          ["Endpoint","Measure","n/events","Unadjusted OR/HR (95% CI)","P","BH FDR",
           "Subtype adjusted OR/HR (95% CI)","P","BH FDR",
           "Subtype + age + sex + WBC OR/HR (95% CI)","P","BH FDR"],s1,
          "Effect per one-unit higher ZEB balance. BH correction across six endpoints within each model. EFS/OS P and FDR are retained as audit outputs only: fully adjusted Cox fits had sparse-subtype convergence warnings and do not support survival inference (see S4).",
          audit_survival=True)

dataset = read(DATA/"supplementary_tables/Table_S2_dataset_units.tsv")
for r in dataset:
    r["dataset"] = r["dataset"].replace("P枚l枚nen", "Pölönen")
figure_map = {
    "GSE142522": "F1; S1",
    "GSE195812": "F1, F3A; S1",
    "GSE206710": "F1, F3A; S1",
    "Park/HTA": "F1; S1",
    "Pölönen": "F2–F4, F7; S2, S8",
    "GSE146901": "F5; S4, S5, S9",
    "Lim": "F6; S3, S8",
    "GSE162280": "F4; S6",
    "HiChIP": "S5, S9",
    "scATAC": "S5, S9",
}
rows2 = [[r["dataset"],r["modality"],r["biological_unit"],r["count"],r["role"],
          next((v for k,v in figure_map.items() if k in r["dataset"]), "see text")]
         for r in dataset]
add_sheet("Table S2","Supplementary Table S2 | Data resources and biological units",
          ["Resource","Context / assay","Independent unit","N / eligible units","Evidentiary role","Figure"], rows2,
          "Counts reflect each analysis eligibility rule; cells, libraries, patients and pooled calls are distinct units.")

subtypes = read(DATA/"supplementary_tables/Table_S4_full_subtype_residuals.tsv")
s3=[]
for r in subtypes:
    s3.append([r["subtype"],int(r["n"]),fnum(r["balance_median"]),fnum(r["dev_median"]),
               ci(r,"residual_within_median","residual_within_ci_low","residual_within_ci_high"),
               ci(r,"adjusted_effect","adjusted_effect_ci_low","adjusted_effect_ci_high"),
               pnum(r["adjusted_effect_hc3_p"]),pnum(r["adjusted_effect_bh_fdr"]),
               f'{r["n_outside_normal_dev_range"]}/{r["n"]}',
               fnum(r["normal_reference_median_sensitivity"])])
add_sheet("Table S3","Supplementary Table S3 | All 17 subtype developmental residuals",
          ["Subtype","n","Raw balance median","Developmental coordinate median",
           "Within-cohort residual median (bootstrap 95% CI)",
           "Adjusted subtype effect (bootstrap 95% CI)","HC3 P","BH FDR",
           "Outside normal range","Normal-reference residual median"],s3,
          "Primary residual is within-cohort observed minus fitted. Bootstrap: 3,000 stratified patient replicates. Normal projection has a separately standardized scale.")

add_sheet("Table S4","Supplementary Table S4 | Fully adjusted clinical sensitivity",
          ["Endpoint","Method","n/events","OR/HR (95% CI)","P","BH FDR",
           "Balance PH P","Global PH P","Fit status"],s4,
          "Model includes 17-subtype indicators, age, sex and log10 white-cell count. EFS/OS Cox fits had sparse-subtype convergence warnings: their P/FDR values are audit outputs only, not evidence of survival association. PH tests do not resolve convergence instability.",
          audit_survival=True)

q1 = read(GATE/"q1_baseline.tsv")
q1adj = read(GATE/"q1_sensitivity_ols.tsv")
q2 = read(GATE/"q2_longitudinal.tsv")
q3 = read(GATE/"q3_state.tsv")
metric_names={"ZEB1_logcpm":"ZEB1","ZEB2_logcpm":"ZEB2","LMO2_logcpm":"LMO2","balance":"ZEB balance"}
rows5=[]
for metric in metric_names:
    b = next(r for r in q1 if r["cohort"]=="combined" and r["metric"]==metric)
    d = next(r for r in q1 if r["cohort"]=="discovery" and r["metric"]==metric)
    e = next(r for r in q1 if r["cohort"]=="extension" and r["metric"]==metric)
    st = next(r for r in q3 if r["group"]=="all_day0_pos_minus_neg" and r["metric"]==metric)
    day = next(r for r in q2 if r["group"]=="all_paired" and r["metric"]==metric)
    adj = [r for r in q1adj if r["y"]==metric]
    fmt_d=lambda r,k,lo,hi:f'{fnum(r[k])} ({fnum(r[lo])}–{fnum(r[hi])})'
    rows5.append([metric_names[metric],f'{b["n_a"]}/{b["n_b"]}',
                  fmt_d(b,"diff_a_minus_b","boot_ci95_lo","boot_ci95_hi"),
                  fnum(b["hedges_g_a_minus_b"]),pnum(b["welch_p"]),
                  fmt_d(d,"diff_a_minus_b","boot_ci95_lo","boot_ci95_hi"),
                  fmt_d(e,"diff_a_minus_b","boot_ci95_lo","boot_ci95_hi"),
                  fmt_d(st,"mean","boot_ci95_lo","boot_ci95_hi"),pnum(st["wilcoxon_p"]),
                  fmt_d(day,"mean","boot_ci95_lo","boot_ci95_hi"),pnum(day["wilcoxon_p"]),
                  "; ".join(f'{r["model"]}: {fnum(r["coef_refractory"])} [{pnum(r["p"])}]' for r in adj)])
add_sheet("Table S5","Supplementary Table S5 | Lim Gate A patient-level results",
          ["Metric","Day0 IF/response n","Day0 IF–response mean diff (bootstrap 95% CI)",
           "Hedges g","Welch P","Discovery mean diff (95% CI)","Extension mean diff (95% CI)",
           "ZBTB16 pos–neg paired mean diff (95% CI)","Paired Wilcoxon P",
           "Day28–Day0 paired mean diff (95% CI)","Paired Wilcoxon P",
           "Day0 one-covariate sensitivity: coefficient [P]"], rows5,
          "CIs shown are bootstrap intervals for mean differences, not Hedges g. Day0 n=54; state pairs n=41; Day0/Day28 pairs n=12. Models add one covariate block separately.")

cases=read(DATA/"S11/S11_full_case_manifest_and_counts.tsv")
rows6=[[r["case"],r["upn"],r["partner"],r["phenotype"],int(r["library_counts"]),
        int(r["ZEB1"]),int(r["ZEB2"]),fnum(r["ZEB1_cpm"]),fnum(r["ZEB2_cpm"]),
        fnum(r["log2_ZEB1_over_ZEB2"])] for r in cases]
add_sheet("Table S6","Supplementary Table S6 | Complete BCL11B-rearranged case manifest",
          ["Case","UPN","Rearrangement partner","Phenotype","Assigned library counts",
           "ZEB1 counts","ZEB2 counts","ZEB1 CPM","ZEB2 CPM","log2(ZEB1/ZEB2)"],rows6,
          "GSE162280 lesion-defined series; all 12 cases retained. CPM and ratios are source-table values. No ordinary T-ALL control group is available.")

hic=read(DATA/"S10/S10A-C_HiChIP_sample_contacts.tsv")
rows7=[[r["sample"],r["group"],int(r["n_contacts"]),int(r["ZEB1"]),int(r["ZEB2"]),
        fnum(r["ZEB1_per_M"]),fnum(r["ZEB2_per_M"]),fnum(r["ZEB2_over_ZEB1"]),
        fnum(r["ZEB1_prom_per_M"]),fnum(r["ZEB2_prom_per_M"])] for r in hic]
add_sheet("Table S7 HiChIP","Supplementary Table S7A | HiChIP sample manifest",
          ["Sample","Group","Filtered contacts","ZEB1 contacts","ZEB2 contacts",
           "ZEB1 per million","ZEB2 per million","ZEB2/ZEB1 ratio",
           "ZEB1 promoter per million","ZEB2 promoter per million"],rows7,
          "Ten usable contact libraries. Gene-body ratios are locus-associated contacts, not transcriptional effects.")
atac=read(DATA/"S10/S10D_scATAC_peak_sets.tsv")
rows7b=[[r["sample"],r["group"],r["source"],int(r["n_peaks"]),int(r["ZEB1_peaks"]),
         int(r["ZEB2_peaks"]),int(r["ZEB1_prom_peaks"]),int(r["ZEB2_prom_peaks"]),
         fnum(r["ZEB1_per_10k_peaks"]),fnum(r["ZEB2_per_10k_peaks"])] for r in atac]
add_sheet("Table S7 scATAC","Supplementary Table S7B | scATAC peak-set manifest",
          ["Sample","Group","Source","Peaks","ZEB1 peaks","ZEB2 peaks",
           "ZEB1 promoter peaks","ZEB2 promoter peaks","ZEB1 per 10k","ZEB2 per 10k"],rows7b,
          "Seven peak sets, including primary and PDX sources. This assay does not independently reproduce the HiChIP ratio.")

ROBUST = DATA / "revision3_robustness"
crossfit = read(ROBUST / "crossfit_subtype_summary.tsv")
add_sheet("Table S8 Crossfit", "Supplementary Table S8A | Leave-one-subtype-out residuals",
          ["Subtype", "n", "Primary median residual", "Cross-fitted median residual",
           "Outside training coordinate range"],
          [[r["subtype"], int(r["n"]), fnum(r["median_primary"]),
            fnum(r["median_crossfit"]), int(r["n_out_of_training_range"])]
           for r in crossfit],
          "Each tested subtype was excluded from spline training; all patients in that subtype were then predicted. This is cohort-internal cross-fitting, not external validation.")
alternatives = read(ROBUST / "alternative_coordinates_subtype_summary.tsv")
add_sheet("Table S8 Scores", "Supplementary Table S8B | Alternative developmental-expression scores",
          ["Definition", "Subtype", "n", "Median residual", "Negative rank / 17", "Development-only R2"],
          [[r["definition"],r["subtype"],int(r["n"]),fnum(r["median_residual"]),
            int(r["rank_negative"]),fnum(r["model_r2"],3)] for r in alternatives],
          "Leave-one-marker-out scores and an expanded normal-direction marker set were fitted within the same diagnostic cohort. Expanded genes have only two normal-source directional checks; all are expression proxies.")
models = read(ROBUST / "model_comparison.tsv")
add_sheet("Table S8 Models", "Supplementary Table S8C | Comparable model performance",
          ["Model", "N in full model", "Model df", "Raw R2", "Adjusted R2",
           "Mean five-fold CV R2", "CV range across five repeats", "Incremental R2 after development", "Partial R2"],
          [[r["model"],int(r["n"]),int(r["df_model"]),fnum(r["raw_r2"],3),
            fnum(r["adjusted_r2"],3),fnum(r["cv_r2_mean"],3),
            f'{fnum(r["cv_r2_min"],3)}–{fnum(r["cv_r2_max"],3)}',
            fnum(r["incremental_r2_after_development"],3),
            fnum(r["partial_r2_subtype_given_development"],3)] for r in models],
          "The same 1,309 patients enter full fits; 1 singleton author-IP 'Other' sample is held in training for all CV models, leaving 1,308 held-out predictions per repeat. No patient contributes to both train and test within a fold.")
contrasts = read(ROBUST / "bcl11b_pairwise_bootstrap.tsv")
matched = read(ROBUST / "bcl11b_etp_matched_summary.tsv")[0]
add_sheet("Table S8 Contrasts", "Supplementary Table S8D | BCL11B focused contrasts",
          ["Comparison", "n BCL11B", "n comparator", "Median-residual difference (95% CI)",
           "Mean matched difference (95% CI)", "Median / maximum coordinate distance"],
          [[r["comparison"],int(r["n_BCL11B"]),int(r["n_comparator"]),
            f'{fnum(r["median_difference"])} ({fnum(r["ci_low"])}–{fnum(r["ci_high"])})',
            "—","—"] for r in contrasts] +
          [["BCL11B vs nearest ETP-like",int(matched["n_pairs"]),int(matched["n_pairs"]),"—",
            f'{fnum(matched["mean_residual_difference"])} ({fnum(matched["ci_low"])}–{fnum(matched["ci_high"])})',
            f'{fnum(matched["median_distance"],3)} / {fnum(matched["max_distance"],3)}']],
          "Three fixed pairwise median comparisons use 3,000 within-group patient bootstraps. The ETP match is deterministic one-to-one nearest-coordinate matching without replacement; matched-pair mean CI uses 3,000 pair bootstraps. No causal interpretation.")

OUT.parent.mkdir(parents=True,exist_ok=True)
book.save(OUT)

# Main Table 1: concise, manuscript-ready dataset/unit/evidence map.
main_rows=[]
for r in dataset:
    main_rows.append([r["dataset"],r["modality"],r["biological_unit"],r["count"],r["role"]])
def esc(s): return str(s).replace("|","/")
md=["**Table 1. Data resources, independent units and evidentiary roles.**",
    "",
    "| Resource | Context | Independent unit | N / eligible units | Main evidentiary role |",
    "|---|---|---|---|---|"]
md += ["| "+" | ".join(esc(x) for x in row)+" |" for row in main_rows]
md += ["","Counts refer to the eligible biological units used in each analysis. The ten HiChIP libraries and seven scATAC peak sets are distinct assays; pooled loops are not patient replicates."]
(MAIN/"Table_1_Datasets_and_Analysis_Units.md").write_text("\n".join(md)+"\n",encoding="utf-8")
print(OUT)
print(MAIN/"Table_1_Datasets_and_Analysis_Units.md")
