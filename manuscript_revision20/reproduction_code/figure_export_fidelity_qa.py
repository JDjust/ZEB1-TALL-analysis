"""Verify current editable figure images, clipping and line types against PDFs."""
from pathlib import Path
import importlib.util,json,hashlib
from zipfile import ZipFile
from io import BytesIO
import fitz
from pptx import Presentation
from PIL import Image
from lxml import etree

B=Path(__file__).resolve().parent.parent;F=B/'figures';Q=B/'QA'
names=[f'Fig{i}' for i in range(1,7)]+[f'SFig{i}' for i in range(1,7)]
spec=importlib.util.spec_from_file_location('current_export',Path(__file__).with_name('export_editable_ppt.py'))
module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
reports=[]
for name in names:
    pp=F/name/(name+'.pptx');pdf=F/name/(name+'.pdf')
    with ZipFile(pp) as z:
        assert z.testzip() is None
        slide=etree.fromstring(z.read('ppt/slides/slide1.xml'))
    pres=Presentation(pp);paths=[s for s in pres.slides[0].shapes if s.name=='Editable plot path']
    pictures=[s for s in pres.slides[0].shapes if s.shape_type==13]
    with fitz.open(pdf) as doc:
        page=doc[0];blocks=[b for b in page.get_text('dict')['blocks'] if b['type']==1]
        drawings,clip_info=module.source_clipped_drawings(page)
        ordered,overlay_seqnos=module.source_paint_events(page,drawings,page.get_text('dict')['blocks'])
    if ordered is not None:
        assert len(ordered)==len(pres.slides[0].shapes)
        for (seqno,kind,item),shape in zip(ordered,pres.slides[0].shapes):
            if kind=='draw':assert shape.name=='Editable plot path'
            elif kind=='text':assert shape.text==item[1]['text']
            else:assert shape.shape_type==13
    assert len(blocks)==len(pictures)
    soft=[]
    for number,(block,picture) in enumerate(zip(blocks,pictures)):
        expected=module.image_with_original_alpha(block)
        assert hashlib.sha256(expected).digest()==hashlib.sha256(picture.image.blob).digest()
        if block.get('mask'):
            im=Image.open(BytesIO(picture.image.blob))
            assert 'A' in im.getbands()
            extrema=im.getchannel('A').getextrema()
            assert extrema[0]<extrema[1]
            soft.append({'picture_index':number,'pixels':im.size,'alpha_range':extrema,'exact_original_rgb_and_mask':True})
    assert len(drawings)==len(paths)
    ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
    actual_dash=len(slide.findall('.//a:custDash',ns))
    expected_dash=sum(bool(module.re.match(r'\[([^\]]+)\]',d.get('dashes','') or '')) for d in drawings)
    assert actual_dash==expected_dash
    # Every emitted native path keeps the recomputed visible PDF geometry. The
    # source clipping helper rejects unsupported crossings rather than omitting.
    for shape,drawing in zip(paths,sorted(drawings,key=lambda d:d.get('seqno',0))):
        r=drawing['rect']
        assert abs(shape.left/12700-r.x0)<.001
        assert abs(shape.top/12700-r.y0)<.001
        assert abs(shape.width/12700-max(r.width,.01))<.001
        assert abs(shape.height/12700-max(r.height,.01))<.001
    reports.append(dict(figure=name,status='PASS',pdf_sha256=sha(pdf),pptx_sha256=sha(pp),raster_layers=len(blocks),original_soft_masks=soft,native_paths=len(paths),rectangular_clipping=clip_info,native_custom_dash_paths=actual_dash,source_raster_overlay_paint_order={'ordered':ordered is not None,'white_overlay_seqnos':overlay_seqnos,'all_native_path_image_text_event_kinds_match_original_pdf':ordered is not None}))
    print(name,'EXPORT_FIDELITY_PASS',flush=True)
(Q/'EXPORT_FIDELITY_20261012.json').write_text(json.dumps({'status':'PASS','basis':'Current PPTX ZIP and native geometry/image bytes checked against the current authoritative PDF, not just receipt declarations','figures':reports},indent=2),encoding='utf-8')
