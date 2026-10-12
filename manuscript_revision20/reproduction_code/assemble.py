from pathlib import Path
import importlib.util, sys, json
import fitz, pandas as pd
from PIL import Image,ImageDraw
from pptx import Presentation
B=Path(__file__).resolve().parent.parent;F=B/'figures';Q=B/'QA';Q.mkdir(exist_ok=True)
names=sys.argv[1:] or [f'Fig{i}' for i in range(1,7)]+[f'SFig{i}' for i in range(1,7)]
spec=importlib.util.spec_from_file_location('native',Path(__file__).with_name('export_editable_ppt.py'));mod=importlib.util.module_from_spec(spec);spec.loader.exec_module(mod);mod.OUT=F
reports=[]
for n in names:
    p=F/n/(n+'.pdf');d=fitz.open(p);page=d[0]
    spans=[s for b in page.get_text('dict')['blocks'] if b['type']==0 for l in b['lines'] for s in l['spans']]
    outside=[s['text'] for s in spans if s['bbox'][0]<-.5 or s['bbox'][1]<-.5 or s['bbox'][2]>page.rect.width+.5 or s['bbox'][3]>page.rect.height+.5]
    reports.append(dict(figure=n,pages=len(d),width_mm=page.rect.width/72*25.4,height_mm=page.rect.height/72*25.4,ratio=page.rect.width/page.rect.height,min_font_pt=min(s['size'] for s in spans),nonblack=[s['text'] for s in spans if s['color']!=0],outside_page_text=outside))
    d.close();mod.export(n)
    pr=Presentation(F/n/(n+'.pptx'));edited=False
    for shape in pr.slides[0].shapes:
        if shape.has_text_frame and shape.text.strip()==n.lower().replace('fig',''):
            continue
        if shape.has_text_frame and shape.text.strip() in list('abcdefghijk'):
            shape.text_frame.paragraphs[0].runs[0].text='EDIT';edited=True;break
    assert edited,n
    pr.save(Q/(n+'_edit_probe.pptx'))
p=Q/'PDF_GEOMETRY_QA.json'
if p.exists():old=json.loads(p.read_text());reports=[r for r in old if r['figure'] not in names and r['figure'] in [f'Fig{i}' for i in range(1,7)]+[f'SFig{i}' for i in range(1,7)]]+reports
p.write_text(json.dumps(reports,indent=2))
print(json.dumps([{k:r[k] for k in ['figure','ratio','min_font_pt','outside_page_text']} for r in reports]))

allnames=[f'Fig{i}' for i in range(1,7)]+[f'SFig{i}' for i in range(1,7)]
if all((F/n/(n+'.pdf')).exists() for n in allnames):
    for label,ns in [('MAIN',allnames[:6]),('SUPPLEMENTARY',allnames[6:]),('COMPLETE',allnames)]:
        out=fitz.open()
        for n in ns:
            with fitz.open(F/n/(n+'.pdf')) as d:out.insert_pdf(d)
        target=B/('ZEB1_'+label+'_FIGURES.pdf')
        if target.exists():target.unlink()
        out.save(target);out.close()
        cols=3;tw=700;th=510;sheet=Image.new('RGB',(tw*cols,th*((len(ns)+cols-1)//cols)),'#EEEEEE');draw=ImageDraw.Draw(sheet)
        for j,n in enumerate(ns):
            im=Image.open(F/n/(n+'.png')).convert('RGB');im.thumbnail((tw-20,th-40));x=(j%cols)*tw+(tw-im.width)//2;y=(j//cols)*th+30;sheet.paste(im,(x,y));draw.text(((j%cols)*tw+12,(j//cols)*th+9),n,fill='black')
        sheet.save(B/('ZEB1_'+label+'_CONTACT_SHEET.png'))
